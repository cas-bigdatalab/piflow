from __future__ import annotations

import argparse
import csv
import io
import json
import math
import re
import statistics
import tarfile
import unicodedata
import zipfile
import posixpath
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Iterator


_SUPPORTED_SUFFIXES = {".csv", ".tsv", ".xlsx", ".json", ".jsonl"}
_FORMAT_PRIORITY = {"csv": 0, "xlsx": 1, "tsv": 2, "jsonl": 3, "json": 4}
_MAX_MEMBER_BYTES = 256 * 1024 * 1024
_NUMBER_RE = re.compile(
    r"^[\s~≈<>≤≥]*(?P<number>[+-]?(?:(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?|\.\d+)(?:[eE][+-]?\d+)?)\s*(?P<unit>.*)$"
)
_RANGE_RE = re.compile(r"^\s*[+-]?(?:\d+(?:\.\d+)?|\.\d+)\s*[-–—~]\s*\d")

_METRIC_ALIASES: dict[str, set[str]] = {
    "specific_surface_area": {
        "specific_surface_area",
        "specific surface area",
        "surface_area",
        "surface area",
        "bet_surface_area",
        "bet surface area",
        "bet specific surface area",
        "BET比表面积",
        "比表面积",
    },
    "capacity": {
        "capacity",
        "discharge_capacity",
        "discharge capacity",
        "specific_capacity",
        "specific capacity",
        "放电容量",
        "比容量",
        "容量",
    },
}


@dataclass(frozen=True)
class ParsedMetric:
    value: float
    assumed_unit: bool


@dataclass
class Candidate:
    name: str
    format_name: str
    values: list[float]
    ignored_count: int
    assumed_unit_count: int
    available_fields: set[str]


def summarize_numeric_metric(
    input_path: str | Path,
    output_path: str | Path,
    *,
    dataset_label: str,
    metric: str,
    order: str = "desc",
    top_n: int = 30,
    target_unit: str = "m2/g",
) -> dict[str, Any]:
    source = Path(input_path).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(f"input artifact not found: {source}")

    normalized_order = str(order).strip().lower()
    if normalized_order not in {"asc", "desc"}:
        raise ValueError("order must be asc or desc")
    normalized_top_n = int(top_n)
    if normalized_top_n <= 0:
        raise ValueError("top_n must be greater than 0")

    metric_key = _resolve_metric_key(metric)
    normalized_target_unit = _normalize_target_unit(target_unit)
    aliases = _METRIC_ALIASES.get(metric_key, {metric_key})
    candidates: list[Candidate] = []
    scanned_files: list[str] = []
    discovered_fields: set[str] = set()

    for artifact_name, payload in _iter_supported_artifacts(source):
        scanned_files.append(artifact_name)
        for candidate_name, format_name, rows in _record_groups(artifact_name, payload):
            candidate = _evaluate_candidate(
                candidate_name,
                format_name,
                rows,
                aliases=aliases,
                metric_key=metric_key,
                target_unit=normalized_target_unit,
            )
            discovered_fields.update(candidate.available_fields)
            if candidate.values:
                candidates.append(candidate)

    if not candidates:
        raise ValueError(
            "METRIC_FIELD_NOT_FOUND: "
            f"dataset={dataset_label!r}, metric={metric_key!r}, "
            f"scanned_files={scanned_files}, "
            f"available_fields={sorted(discovered_fields)[:100]}"
        )

    chosen = min(
        candidates,
        key=lambda item: (
            -len(item.values),
            _FORMAT_PRIORITY.get(item.format_name, 99),
            item.name.casefold(),
        ),
    )
    selected = sorted(chosen.values, reverse=normalized_order == "desc")[:normalized_top_n]
    result = {
        "contract_version": "metric_summary_v1",
        "dataset_name": str(dataset_label).strip(),
        "metric_key": metric_key,
        "selected_count": len(selected),
        "available_count": len(chosen.values),
        "mean_value": statistics.fmean(selected),
        "min_value": min(selected),
        "max_value": max(selected),
        "unit": normalized_target_unit,
        "source_table": chosen.name,
        "ignored_value_count": chosen.ignored_count,
        "assumed_unit_count": chosen.assumed_unit_count,
    }
    _write_csv_row(Path(output_path), result)
    return result


def _iter_supported_artifacts(source: Path) -> Iterator[tuple[str, bytes]]:
    if source.is_dir():
        for candidate in sorted(path for path in source.rglob("*") if path.is_file()):
            if candidate.suffix.lower() in _SUPPORTED_SUFFIXES:
                yield candidate.relative_to(source).as_posix(), _read_limited(candidate)
        return

    if tarfile.is_tarfile(source):
        with tarfile.open(source, mode="r:*") as archive:
            for member in archive:
                suffix = Path(member.name).suffix.lower()
                if not member.isfile() or suffix not in _SUPPORTED_SUFFIXES:
                    continue
                if member.size > _MAX_MEMBER_BYTES:
                    continue
                handle = archive.extractfile(member)
                if handle is not None:
                    yield member.name, handle.read(_MAX_MEMBER_BYTES + 1)
        return

    # XLSX is itself a ZIP container.  Treat supported table files as the
    # actual input before applying generic ZIP-archive discovery.
    if source.suffix.lower() in _SUPPORTED_SUFFIXES:
        yield source.name, _read_limited(source)
        return

    if zipfile.is_zipfile(source):
        with zipfile.ZipFile(source) as archive:
            for info in archive.infolist():
                if info.is_dir() or Path(info.filename).suffix.lower() not in _SUPPORTED_SUFFIXES:
                    continue
                if info.file_size > _MAX_MEMBER_BYTES:
                    continue
                yield info.filename, archive.read(info)
        return
    raise ValueError(f"no supported structured data found in input artifact: {source}")


def _read_limited(path: Path) -> bytes:
    if path.stat().st_size > _MAX_MEMBER_BYTES:
        raise ValueError(f"structured data file is too large: {path}")
    return path.read_bytes()


def _record_groups(name: str, payload: bytes) -> Iterator[tuple[str, str, Iterable[dict[str, Any]]]]:
    suffix = Path(name).suffix.lower()
    format_name = suffix.lstrip(".")
    if suffix in {".csv", ".tsv"}:
        delimiter = "," if suffix == ".csv" else "\t"
        yield name, format_name, _read_delimited(payload, delimiter)
        return
    if suffix == ".xlsx":
        yield from _read_xlsx_groups(name, payload)
        return
    if suffix == ".jsonl":
        yield name, format_name, _read_jsonl(payload)
        return
    if suffix == ".json":
        yield name, format_name, _read_json(payload)


def _decode_text(payload: bytes) -> str:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"unable to decode structured text: {last_error}")


def _read_delimited(payload: bytes, delimiter: str) -> list[dict[str, Any]]:
    reader = csv.DictReader(io.StringIO(_decode_text(payload)), delimiter=delimiter)
    return [dict(row) for row in reader if row]


def _read_jsonl(payload: bytes) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for line_number, line in enumerate(_decode_text(payload).splitlines(), start=1):
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"invalid JSONL at line {line_number}: {exc.msg}") from exc
        records.extend(_expand_json_rows(value))
    return records


def _read_json(payload: bytes) -> list[dict[str, Any]]:
    try:
        value = json.loads(_decode_text(payload))
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid JSON: {exc.msg}") from exc
    return _expand_json_rows(value)


def _expand_json_rows(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    if not isinstance(value, dict):
        return []
    for wrapper in ("records", "items", "data", "rows"):
        wrapped = value.get(wrapper)
        if isinstance(wrapped, list):
            return [item for item in wrapped if isinstance(item, dict)]
    annotation = value.get("annotation")
    if isinstance(annotation, list) and annotation:
        return [item for item in annotation if isinstance(item, dict)]
    return [value]


def _read_xlsx_groups(name: str, payload: bytes) -> Iterator[tuple[str, str, list[dict[str, Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        yield from _read_xlsx_groups_stdlib(name, payload)
        return

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            raw_rows = list(sheet.iter_rows(values_only=True))
            header_index = next(
                (index for index, row in enumerate(raw_rows) if any(cell not in (None, "") for cell in row)),
                None,
            )
            if header_index is None:
                continue
            headers = _unique_headers(raw_rows[header_index])
            records = [
                {headers[index]: value for index, value in enumerate(row[: len(headers)])}
                for row in raw_rows[header_index + 1 :]
                if any(value not in (None, "") for value in row)
            ]
            yield f"{name}#{sheet.title}", "xlsx", records
    finally:
        workbook.close()


def _read_xlsx_groups_stdlib(
    name: str,
    payload: bytes,
) -> Iterator[tuple[str, str, list[dict[str, Any]]]]:
    """Read ordinary XLSX tables without third-party dependencies.

    This fallback intentionally reads cached cell values only.  It is enough for
    exported scientific data tables and keeps execution nodes free from an
    openpyxl deployment requirement.
    """
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ns = {"x": spreadsheet_ns, "r": office_rel_ns}

    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = set(archive.namelist())
        if "xl/workbook.xml" not in names or "xl/_rels/workbook.xml.rels" not in names:
            raise ValueError(f"invalid XLSX workbook structure: {name}")

        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in names:
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in shared_root.findall(f"{{{spreadsheet_ns}}}si"):
                shared_strings.append(
                    "".join(node.text or "" for node in item.iter(f"{{{spreadsheet_ns}}}t"))
                )

        relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relationship_targets = {
            rel.attrib.get("Id", ""): rel.attrib.get("Target", "")
            for rel in relationships_root.findall(f"{{{package_rel_ns}}}Relationship")
        }
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in workbook_root.findall("x:sheets/x:sheet", ns):
            sheet_name = sheet.attrib.get("name", "sheet")
            relationship_id = sheet.attrib.get(f"{{{office_rel_ns}}}id", "")
            target = relationship_targets.get(relationship_id, "")
            if not target:
                continue
            if target.startswith("/"):
                worksheet_path = target.lstrip("/")
            else:
                worksheet_path = posixpath.normpath(posixpath.join("xl", target))
            if worksheet_path not in names:
                continue

            worksheet_root = ET.fromstring(archive.read(worksheet_path))
            raw_rows: list[list[Any]] = []
            for row in worksheet_root.findall(".//x:sheetData/x:row", ns):
                values: dict[int, Any] = {}
                for cell in row.findall("x:c", ns):
                    reference = cell.attrib.get("r", "")
                    column_index = _xlsx_column_index(reference)
                    values[column_index] = _xlsx_cell_value(
                        cell,
                        shared_strings=shared_strings,
                        spreadsheet_ns=spreadsheet_ns,
                    )
                if values:
                    width = max(values) + 1
                    raw_rows.append([values.get(index) for index in range(width)])

            header_index = next(
                (index for index, row in enumerate(raw_rows) if any(cell not in (None, "") for cell in row)),
                None,
            )
            if header_index is None:
                continue
            headers = _unique_headers(raw_rows[header_index])
            records = [
                {
                    headers[index]: row[index] if index < len(row) else None
                    for index in range(len(headers))
                }
                for row in raw_rows[header_index + 1 :]
                if any(value not in (None, "") for value in row)
            ]
            yield f"{name}#{sheet_name}", "xlsx", records


def _xlsx_column_index(reference: str) -> int:
    letters = "".join(character for character in reference if character.isalpha()).upper()
    if not letters:
        return 0
    result = 0
    for character in letters:
        result = result * 26 + (ord(character) - ord("A") + 1)
    return result - 1


def _xlsx_cell_value(
    cell: ET.Element,
    *,
    shared_strings: list[str],
    spreadsheet_ns: str,
) -> Any:
    cell_type = cell.attrib.get("t", "")
    if cell_type == "inlineStr":
        inline = cell.find(f"{{{spreadsheet_ns}}}is")
        if inline is None:
            return ""
        return "".join(node.text or "" for node in inline.iter(f"{{{spreadsheet_ns}}}t"))

    value_node = cell.find(f"{{{spreadsheet_ns}}}v")
    if value_node is None or value_node.text is None:
        return None
    raw = value_node.text
    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except (ValueError, IndexError):
            return raw
    if cell_type in {"str", "e"}:
        return raw
    if cell_type == "b":
        return raw == "1"
    try:
        number = float(raw)
    except ValueError:
        return raw
    return int(number) if number.is_integer() else number


def _unique_headers(row: Iterable[Any]) -> list[str]:
    headers: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(row, start=1):
        base = str(value).strip() if value not in (None, "") else f"column_{index}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _evaluate_candidate(
    name: str,
    format_name: str,
    rows: Iterable[dict[str, Any]],
    *,
    aliases: set[str],
    metric_key: str,
    target_unit: str,
) -> Candidate:
    values: list[float] = []
    ignored = 0
    assumed = 0
    fields: set[str] = set()
    for row in rows:
        fields.update(_collect_field_names(row))
        match = _best_metric_match(row, aliases)
        if match is None:
            continue
        raw_value, unit_hint = match
        parsed = _parse_metric_value(
            raw_value,
            unit_hint=unit_hint,
            metric_key=metric_key,
            target_unit=target_unit,
        )
        if parsed is None:
            ignored += 1
            continue
        values.append(parsed.value)
        assumed += int(parsed.assumed_unit)
    return Candidate(name, format_name, values, ignored, assumed, fields)


def _collect_field_names(value: Any, *, depth: int = 0) -> set[str]:
    if depth > 5:
        return set()
    fields: set[str] = set()
    if isinstance(value, dict):
        for key, nested in value.items():
            fields.add(str(key))
            fields.update(_collect_field_names(nested, depth=depth + 1))
    elif isinstance(value, list):
        for nested in value[:20]:
            fields.update(_collect_field_names(nested, depth=depth + 1))
    return fields


def _best_metric_match(row: dict[str, Any], aliases: set[str]) -> tuple[Any, str] | None:
    matches: list[tuple[int, Any, str]] = []

    def visit(value: Any, depth: int = 0) -> None:
        if depth > 8:
            return
        if isinstance(value, dict):
            logical_name = _first_present(
                value,
                ("name", "property_name", "propertyName", "field", "attribute", "label", "key"),
            )
            logical_value = _first_present(
                value,
                ("value", "property_value", "propertyValue", "raw_value", "content"),
            )
            if logical_name is not None and logical_value is not None:
                score = _field_match_score(str(logical_name), aliases)
                if score:
                    matches.append((score + 10, logical_value, str(value.get("unit") or "")))

            for key, nested in value.items():
                key_text = str(key)
                if key_text.casefold().endswith(("_unit", " unit")):
                    continue
                score = _field_match_score(key_text, aliases)
                if score and not isinstance(nested, (dict, list)):
                    unit = _unit_from_header(key_text)
                    sibling_unit = value.get(f"{key_text}_unit") or value.get(f"{key_text} unit")
                    matches.append((score, nested, str(sibling_unit or unit)))
                if isinstance(nested, (dict, list)):
                    visit(nested, depth + 1)
        elif isinstance(value, list):
            for nested in value:
                visit(nested, depth + 1)

    visit(row)
    if not matches:
        return None
    _, raw_value, unit_hint = max(matches, key=lambda item: item[0])
    return raw_value, unit_hint


def _field_match_score(field_name: str, aliases: set[str]) -> int:
    canonical = _canonical(field_name)
    best = 0
    for alias in aliases:
        alias_key = _canonical(alias)
        if canonical == alias_key:
            best = max(best, 100 + len(alias_key))
        elif len(alias_key) >= 6 and alias_key in canonical:
            best = max(best, 50 + len(alias_key))
    return best


def _first_present(mapping: dict[str, Any], keys: tuple[str, ...]) -> Any:
    for key in keys:
        if key in mapping and mapping[key] not in (None, ""):
            return mapping[key]
    return None


def _parse_metric_value(
    raw_value: Any,
    *,
    unit_hint: str,
    metric_key: str,
    target_unit: str,
) -> ParsedMetric | None:
    if raw_value is None or isinstance(raw_value, bool):
        return None
    if isinstance(raw_value, dict):
        nested_value = _first_present(raw_value, ("value", "raw_value", "content"))
        if nested_value is None:
            return None
        unit_hint = str(raw_value.get("unit") or unit_hint)
        raw_value = nested_value

    if isinstance(raw_value, (int, float)):
        number = float(raw_value)
        value_unit = ""
    else:
        text = str(raw_value).strip()
        if not text or _RANGE_RE.match(text):
            return None
        matched = _NUMBER_RE.fullmatch(text)
        if matched is None:
            return None
        number = float(matched.group("number").replace(",", ""))
        value_unit = matched.group("unit").strip()
    if not math.isfinite(number):
        return None

    unit = value_unit or unit_hint
    assumed = not bool(str(unit).strip())
    if assumed:
        unit = _default_unit(metric_key)
    converted = _convert_unit(number, str(unit), target_unit)
    if converted is None or not math.isfinite(converted):
        return None
    return ParsedMetric(converted, assumed)


def _unit_from_header(header: str) -> str:
    matched = re.search(r"[\(\[]\s*([^\)\]]+)\s*[\)\]]", header)
    return matched.group(1).strip() if matched else ""


def _default_unit(metric_key: str) -> str:
    if metric_key == "specific_surface_area":
        return "m2/g"
    return ""


def _convert_unit(number: float, source_unit: str, target_unit: str) -> float | None:
    source = _canonical_unit(source_unit)
    target = _canonical_unit(target_unit)
    if source == target:
        return number
    if target == "m2/g":
        factors = {
            "m2/kg": 0.001,
            "cm2/g": 0.0001,
            "cm2/kg": 0.0000001,
        }
        factor = factors.get(source)
        return number * factor if factor is not None else None
    return None


def _normalize_target_unit(unit: str) -> str:
    canonical = _canonical_unit(unit)
    if canonical not in {"m2/g", "m2/kg", "cm2/g", "cm2/kg"}:
        raise ValueError(f"unsupported target_unit: {unit!r}")
    return canonical


def _canonical_unit(unit: str) -> str:
    value = unicodedata.normalize("NFKC", str(unit or "")).casefold().strip()
    value = value.replace("㎡", "m2").replace("²", "2")
    value = value.replace("−", "-").replace("–", "-").replace("—", "-")
    value = value.replace("⁻", "-").replace("¹", "1")
    value = value.replace("·", "").replace("*", "").replace("^", "")
    value = re.sub(r"\s+", "", value)
    value = value.replace("per", "/")
    value = re.sub(r"(m2|cm2)g-?1$", r"\1/g", value)
    value = re.sub(r"(m2|cm2)kg-?1$", r"\1/kg", value)
    return value


def _resolve_metric_key(metric: str) -> str:
    requested = _canonical(metric)
    if not requested:
        raise ValueError("metric must not be empty")
    for key, aliases in _METRIC_ALIASES.items():
        if requested == _canonical(key) or any(requested == _canonical(alias) for alias in aliases):
            return key
    return str(metric).strip()


def _canonical(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(value)).casefold()
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", normalized)


def _write_csv_row(output_path: Path, row: dict[str, Any]) -> None:
    target = output_path.expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(row))
        writer.writeheader()
        writer.writerow({key: _format_value(value) for key, value in row.items()})


def _format_value(value: Any) -> Any:
    if isinstance(value, float):
        return format(value, ".12g")
    return value


def main() -> None:
    parser = argparse.ArgumentParser(description="Discover, rank and summarize one numeric metric.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--dataset_label", required=True)
    parser.add_argument("--metric", required=True)
    parser.add_argument("--order", default="desc")
    parser.add_argument("--top_n", type=int, default=30)
    parser.add_argument("--target_unit", required=True)
    args = parser.parse_args()

    result = summarize_numeric_metric(
        args.input,
        args.output,
        dataset_label=args.dataset_label,
        metric=args.metric,
        order=args.order,
        top_n=args.top_n,
        target_unit=args.target_unit,
    )
    print(
        "[OK] numeric metric summary generated: "
        f"dataset={result['dataset_name']}, selected={result['selected_count']}, "
        f"source={result['source_table']}, output={Path(args.output).resolve()}"
    )


if __name__ == "__main__":
    main()
